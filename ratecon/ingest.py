"""Turn a real document into the raw text the pipeline consumes.

The pipeline is "text in": everything downstream (the prompt, grounding,
projection) works on the plain text of a rate confirmation. This module is the
front door that gets you there from a .txt / .pdf / .docx / image file.

Design notes:

* Each backend is imported lazily and, if missing, raises an actionable error
  naming the exact ``pip install`` (and system binary for OCR). The core
  pipeline and the offline test suite therefore need none of these libraries.
* PDF text is extracted layout-preserved. Rate confirmations are tables, and
  the pipeline's address/date parsing was built against column-aligned text
  (the way the provided fixtures look), so preserving layout matters.
* Images and scanned PDFs go through OCR, which is lossy -- OCR noise is
  exactly what grounding and the confidence engine are there to absorb, so a
  bad scan degrades to low confidence rather than to bad data.
"""

from __future__ import annotations

from pathlib import Path

TEXT_EXT = {".txt", ".text", ".md"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}

SUPPORTED = TEXT_EXT | PDF_EXT | DOCX_EXT | IMAGE_EXT


def _missing(pkg: str, extra: str = "") -> ModuleNotFoundError:
    return ModuleNotFoundError(
        f"reading this file type needs `{pkg}`. Install it with "
        f"`pip install {pkg}`{extra}.")


def to_text(path: str | Path) -> str:
    """Extract the raw text of a document, dispatching on file extension."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="replace")
    if ext in PDF_EXT:
        return _from_pdf(path)
    if ext in DOCX_EXT:
        return _from_docx(path)
    if ext in IMAGE_EXT:
        return _from_image(path)
    # Unknown extension: assume it is already text rather than guessing.
    return path.read_text(encoding="utf-8", errors="replace")


def _from_pdf(path: Path) -> str:
    """PDF -> layout-preserved text. Falls back to OCR for scanned PDFs."""
    try:
        import pdfplumber
    except ModuleNotFoundError:
        raise _missing("pdfplumber")

    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(layout=True) or "")
    text = "\n".join(pages)

    # A born-digital rate con yields real text; a scanned one yields almost
    # nothing here. Fall back to rasterise-then-OCR if the page had no text.
    if len(text.strip()) < 20:
        return _ocr_pdf(path)
    return text


def _ocr_pdf(path: Path) -> str:
    try:
        import pdf2image
        import pytesseract
    except ModuleNotFoundError:
        raise _missing(
            "pdf2image pytesseract",
            " plus the tesseract and poppler binaries "
            "(macOS: `brew install tesseract poppler`)")
    pages = pdf2image.convert_from_path(str(path))
    return "\n".join(pytesseract.image_to_string(p) for p in pages)


def _from_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ModuleNotFoundError:
        raise _missing("python-docx")
    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    # Tables carry the stop/charge grids, so flatten them too.
    for table in document.tables:
        for row in table.rows:
            lines.append("    ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _from_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ModuleNotFoundError:
        raise _missing(
            "pytesseract pillow",
            " plus the tesseract binary (macOS: `brew install tesseract`)")
    return pytesseract.image_to_string(Image.open(str(path)))
